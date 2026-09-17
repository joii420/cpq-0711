"""task-260916 S-UI: generate element-price import fixtures (AC-7 / AC-8 / AC-18 step 1).
Format copied from 素材/元素价格导入模板-用户样例.xlsx: sheet「价格导入」, header 元素符号* / 单价* / 货币 / 计价单位,
price cells are NUMERIC (same as the user sample). Currency / unit left blank (AC-7 wording).
Usage: python3 gen_fixtures.py <out_dir>
"""
import sys, os, hashlib
import openpyxl

FIXTURES = {
    'ac7-import-3rows.xlsx':  [('Cu', 101.13921), ('Zn', 24.123456789), ('Ni', 2.0000000005)],
    'ac8-import-zero.xlsx':   [('Cu', 0.0000000004)],
    'ac18-import-cu.xlsx':    [('Cu', 101.13922)],
}

def main(out):
    os.makedirs(out, exist_ok=True)
    for name, rows in FIXTURES.items():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '价格导入'
        ws.append(['元素符号*', '单价*', '货币', '计价单位'])
        for code, price in rows:
            ws.append([code, price, None, None])
        p = os.path.join(out, name)
        wb.save(p)
        # read back to prove the file holds what we think
        back = [(r[0].value, r[1].value, r[1].data_type) for r in openpyxl.load_workbook(p).active.iter_rows(min_row=2)]
        print(name, hashlib.md5(open(p, 'rb').read()).hexdigest(), back)

if __name__ == '__main__':
    main(sys.argv[1])
