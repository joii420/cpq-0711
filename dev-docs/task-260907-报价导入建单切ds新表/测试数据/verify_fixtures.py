# -*- coding: utf-8 -*-
"""夹具自检：把 test.md §0 的 6 条硬要求逐条断言出来，并打印实际值。

🚨 本脚本自身也遵守 testing.md §4.4：先证明「它能输出东西」——
   每条检查都打印实测数字，全程空输出与「全部通过」长得一模一样。
"""
import glob
import os
import sys

import openpyxl

OUT = os.path.dirname(os.path.abspath(__file__))
SHEETS = ['物料', '客户料号', '物料BOM', '物料与元素BOM', '来料固定加工费', '来料其他费用',
          '来料回收折扣', '自制加工费', '成品其他费用', '组成件其他费用', '组装加工费',
          '组装加工费年降', '电镀费用', '电镀方案', '来料年降', '年降系数']
KNOWN_CUSTOMERS = {'CUST-0001', 'CUST-0002', 'CUST-0004'}   # customer.code 实测可命中
fails = []


def chk(cond, msg):
    print(("  ✅ " if cond else "  ❌ ") + msg)
    if not cond:
        fails.append(msg)


def datarows(ws):
    n = 0
    for r in range(2, ws.max_row + 1):
        if any(ws.cell(r, c).value not in (None, '') for c in range(1, ws.max_column + 1)):
            n += 1
    return n


def colvals(ws, name):
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    i = hdr.index(name) + 1
    return [ws.cell(r, i).value for r in range(2, ws.max_row + 1)
            if ws.cell(r, i).value not in (None, '')]


for path in sorted(glob.glob(os.path.join(OUT, "T260907-*.xlsx"))):
    name = os.path.basename(path)
    wb = openpyxl.load_workbook(path)
    print("\n=== %s ===" % name)
    chk(wb.sheetnames == SHEETS, "16 sheet 齐全且顺序一致（实际 %d 张）" % len(wb.sheetnames))
    counts = {s: datarows(wb[s]) for s in SHEETS}
    print("  行数：" + " ".join("%s=%d" % (k, v) for k, v in counts.items()))

    cust = colvals(wb['客户料号'], '客户编号')
    uniq = sorted(set(cust))
    if '负例' in name:
        if '跨客户' in name:
            chk(len(uniq) == 2,
                "跨客户负例含 2 个客户编号：%s" % uniq)
            chk(set(uniq) <= KNOWN_CUSTOMERS,
                "两个客户编号都在 customer.code 命中（否则先被 D-19 拦住，AC-3 空过）：%s" % uniq)
        else:
            bad = [c for c in uniq if c not in KNOWN_CUSTOMERS]
            chk(len(bad) == 2, "非法客户负例含 2 个非法编号（验 errors 逐条而非只第一条）：%s" % bad)
    elif '空客户料号' in name:
        chk(len(cust) == 0, "客户料号 sheet 只有表头（AC-18）")
        chk(counts['物料'] > 0, "其余 sheet 仍有数据（否则整份是空验证）")
    else:
        chk(uniq == ['CUST-0004'], "主文件单客户（D-23）：%s" % uniq)

    if '空sheet' in name:
        chk(counts['年降系数'] == 0, "年降系数 只有表头（AC-17）")
        chk(sum(v for k, v in counts.items() if k != '年降系数') > 0, "其余 15 张仍有数据")

    if name.startswith('T260907T-主文件'):
        fee = ['来料固定加工费', '来料其他费用', '来料回收折扣', '自制加工费',
               '成品其他费用', '组成件其他费用', '组装加工费', '电镀费用']
        chk(all(counts[s] >= 2 for s in fee),
            "8 张费用表每张 ≥2 行：%s" % {s: counts[s] for s in fee})
        for s in fee:
            v = colvals(wb[s], '销售料号')
            n1 = v.count('T260907T-FG01')
            chk(n1 == 2 and len(v) == 3,
                "%s：FG01=%d 行 / 整表=%d 行（两数不等，漏 WHERE 谓词才抓得住）" % (s, n1, len(v)))
        cp = colvals(wb['客户料号'], '销售料号')
        chk(cp.count('T260907T-FG01') == 2, "AC-21：FG01 有 2 条客户料号行")
        pno = colvals(wb['客户料号'], '客户产品编号')
        chk(len(set(pno)) == len(pno), "客户产品编号互不相同：%s" % pno)
        chk(len(cp) == 3, "期望 lineItemsCount = 3（实际客户料号行数 %d）" % len(cp))
        mats = set(colvals(wb['物料'], '销售料号'))
        chk('T260907T-FG03' in mats and 'T260907T-FG03' not in cp,
            "FG03 有物料无客户料号 —— 守卫「明细行只由客户料号决定」")
        el = set(colvals(wb['物料与元素BOM'], '元素'))
        chk(el <= {'Ag', 'Cu', 'Ni', 'Zn'},
            "元素全部取自 CUST-0004 有行情的集合（AC-8 前置）：%s" % sorted(el))
        rec = set(str(x) for x in colvals(wb['物料与元素BOM'], '材质料号'))
        chk(rec <= {'992', '00006'}, "材质料号取自 material_recipe 实存值：%s" % sorted(rec))
        proc = set(colvals(wb['自制加工费'], '工序编号')) | set(colvals(wb['组装加工费'], '组装工序'))
        chk(proc <= {'Z100', 'Z101'}, "工序编号取自 process_master 实存值：%s" % sorted(proc))

    if '大单量' in name:
        chk(counts['客户料号'] in (200, 1845), "大单量客户料号行数 = %d" % counts['客户料号'])

    # 前缀纪律：所有自造料号必须带前缀（主数据引用列除外）
    strays = [v for v in colvals(wb['物料'], '销售料号') if not str(v).startswith('T260907T-')]
    chk(not strays, "物料表全部料号带 T260907T- 前缀（可一条 SQL 清干净）；例外：%s" % strays[:5])

if fails:
    print("\n❌ 夹具自检失败 %d 条：" % len(fails))
    for f in fails:
        print("   - " + f)
    sys.exit(1)
print("\n✅ 夹具自检全通过")
