#!/usr/bin/env python3
"""Black-box HTTP checks for AC-7 through AC-14.

The runner intentionally knows only the task API contract. It never imports
application classes, touches the database, or deletes test data.
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook


HEADERS = [
    "客户编号",
    "客户名称",
    "客户等级",
    "所属行业",
    "产品分类",
    "所属区域",
    "地址",
    "信用额度",
    "付款方式",
    "备注",
    "联系人姓名",
    "联系人职务",
    "联系人电话",
    "联系人邮箱",
    "联系人微信",
]


class ContractFailure(AssertionError):
    pass


def data_of(response: requests.Response) -> dict[str, Any]:
    try:
        body = response.json()
    except ValueError as exc:
        raise ContractFailure(f"HTTP {response.status_code} returned non-JSON: {response.text[:300]}") from exc
    data = body.get("data")
    return data if isinstance(data, dict) else {}


def check_counts(data: dict[str, Any], *, total: int, customers: int, contacts: int) -> None:
    actual = (data.get("totalRows"), data.get("createdCustomers"), data.get("createdContacts"))
    expected = (total, customers, contacts)
    if actual != expected:
        raise ContractFailure(f"summary mismatch: expected {expected}, got {actual}")


def post_import(session: requests.Session, base: str, fixture: Path) -> requests.Response:
    with fixture.open("rb") as stream:
        return session.post(
            f"{base}/customers/import",
            files={"file": (fixture.name, stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            timeout=30,
        )


def require_any_status(response: requests.Response, statuses: set[int], label: str) -> dict[str, Any]:
    if response.status_code not in statuses:
        raise ContractFailure(f"{label}: expected {sorted(statuses)}, got {response.status_code}: {response.text[:500]}")
    return data_of(response)


def test_template(session: requests.Session, base: str, evidence_dir: Path) -> None:
    # The template endpoint negotiates an xlsx representation. Keep the
    # session's JSON Accept for API responses, but explicitly request xlsx
    # here (some REST stacks answer 406 for Accept: application/json).
    response = session.get(
        f"{base}/customers/import/template",
        headers={"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        timeout=30,
    )
    if response.status_code != 200:
        raise ContractFailure(f"AC-7 template status: expected 200, got {response.status_code}: {response.text[:300]}")
    if not response.content.startswith(b"PK"):
        raise ContractFailure("AC-7 template is not an xlsx zip payload")
    downloaded = evidence_dir / "template-download.xlsx"
    downloaded.write_bytes(response.content)
    wb = load_workbook(downloaded, read_only=True, data_only=True)
    if wb.sheetnames != ["客户"]:
        raise ContractFailure(f"AC-7 sheet mismatch: {wb.sheetnames}")
    ws = wb["客户"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows or list(rows[0]) != HEADERS:
        raise ContractFailure(f"AC-7 header mismatch: {rows[0] if rows else None}")
    if len(rows[0]) != 15:
        raise ContractFailure(f"AC-7 expected 15 columns, got {len(rows[0])}")
    flattened = " ".join(str(cell) for row in rows[:10] for cell in row if cell is not None)
    for forbidden in ("状态", "累计金额", "历史订单数", "平均折扣率"):
        if forbidden in flattened:
            raise ContractFailure(f"AC-7 template contains forbidden field: {forbidden}")
    print("PASS AC-7 template: 客户 sheet, 15 headers, no status/stat fields")


def run_import_cases(session: requests.Session, base: str, fixtures: Path) -> None:
    valid = post_import(session, base, fixtures / "valid_two_rows.xlsx")
    if valid.status_code != 200:
        raise ContractFailure(f"AC-8 expected 200, got {valid.status_code}: {valid.text[:500]}")
    check_counts(data_of(valid), total=2, customers=2, contacts=2)
    print("PASS AC-8 valid_two_rows: 2/2/2")

    conflict = post_import(session, base, fixtures / "existing_code_mix.xlsx")
    conflict_data = require_any_status(conflict, {400, 409}, "AC-9 existing code")
    check_counts(conflict_data, total=2, customers=0, contacts=0)
    if not conflict_data.get("errors"):
        raise ContractFailure("AC-9 response has no conflict errors")
    print("PASS AC-9 existing_code_mix: full-file rejection, 0/0/0")

    duplicate = post_import(session, base, fixtures / "duplicate_in_file.xlsx")
    duplicate_data = require_any_status(duplicate, {400, 409}, "AC-11 file duplicate")
    check_counts(duplicate_data, total=2, customers=0, contacts=0)
    errors = duplicate_data.get("errors") or []
    rows = {item.get("row") for item in errors if isinstance(item, dict)}
    if not ({2, 3} <= rows):
        raise ContractFailure(f"AC-11 expected errors for Excel rows 2 and 3, got {rows}")
    print("PASS AC-11 duplicate_in_file: rows 2+3 reported, 0/0/0")

    multi = post_import(session, base, fixtures / "multiple_errors.xlsx")
    multi_data = require_any_status(multi, {400, 409}, "AC-12 multiple errors")
    check_counts(multi_data, total=2, customers=0, contacts=0)
    multi_errors = multi_data.get("errors") or []
    if len(multi_errors) < 7:
        raise ContractFailure(f"AC-12 expected at least 7 aggregated errors, got {len(multi_errors)}")
    print(f"PASS AC-12 multiple_errors: {len(multi_errors)} errors aggregated, 0/0/0")

    for name in ("not_xlsx.xls", "damaged.xlsx", "wrong_header.xlsx"):
        response = post_import(session, base, fixtures / name)
        require_any_status(response, {400}, f"AC-13 {name}")
        print(f"PASS AC-13 {name}: HTTP 400")

    empty = post_import(session, base, fixtures / "header_only.xlsx")
    if empty.status_code != 200:
        raise ContractFailure(f"AC-14 expected 200, got {empty.status_code}: {empty.text[:500]}")
    check_counts(data_of(empty), total=0, customers=0, contacts=0)
    print("PASS AC-14 header_only: 0/0/0")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--fixtures", type=Path, required=True)
    parser.add_argument("--evidence-dir", type=Path, default=Path("."))
    parser.add_argument("--base-url", default=os.environ.get("CPQ_BASE_URL", "http://localhost:8081/api/cpq"))
    parser.add_argument("--token", default=os.environ.get("CPQ_AUTH_TOKEN"))
    args = parser.parse_args()
    if not args.token:
        print("BLOCKED: CPQ_AUTH_TOKEN is required; no API calls were made", file=sys.stderr)
        return 2
    if not args.fixtures.is_dir():
        print(f"BLOCKED: fixture directory does not exist: {args.fixtures}", file=sys.stderr)
        return 2
    args.evidence_dir.mkdir(parents=True, exist_ok=True)
    session = requests.Session()
    session.headers.update({"Authorization": f"Bearer {args.token}", "Accept": "application/json"})
    try:
        test_template(session, args.base_url.rstrip("/"), args.evidence_dir)
        run_import_cases(session, args.base_url.rstrip("/"), args.fixtures)
    except (requests.RequestException, ContractFailure, OSError) as exc:
        print(f"FAIL/BLOCKED: {exc}", file=sys.stderr)
        return 1
    print("S-API-2 AC-7..AC-14 black-box checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
