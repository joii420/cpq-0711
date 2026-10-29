#!/usr/bin/env python3
"""Render the CPQ whitepaper Markdown to styled HTML and DOCX."""

from __future__ import annotations

import copy
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "CPQ系统白皮书.md"
HTML_OUT = ROOT / "CPQ系统白皮书.html"
DOCX_OUT = ROOT / "CPQ系统白皮书.docx"

NAVY = "092B43"
CYAN = "0BBBD6"
BLUE = "4776E6"
PALE = "EAF7FA"
GRAY = "526D7D"


def render_markdown() -> BeautifulSoup:
    source = SOURCE.read_text(encoding="utf-8")
    html = markdown.markdown(
        source,
        extensions=["tables", "fenced_code", "sane_lists"],
        output_format="html5",
    )
    return BeautifulSoup(html, "html.parser")


def build_html(source_soup: BeautifulSoup) -> None:
    soup = BeautifulSoup("<main class='document'></main>", "html.parser")
    main = soup.main
    nodes = [copy.copy(node) for node in source_soup.contents]
    cover_nodes: list[Tag | NavigableString] = []
    body_nodes: list[Tag | NavigableString] = []
    hit_rule = False
    for node in nodes:
        if isinstance(node, Tag) and node.name == "hr" and not hit_rule:
            hit_rule = True
            continue
        (body_nodes if hit_rule else cover_nodes).append(node)

    cover = soup.new_tag("section", attrs={"class": "cover"})
    brand = soup.new_tag("div", attrs={"class": "brand-mark"})
    brand.string = "CPQ"
    cover.append(brand)
    for node in cover_nodes:
        cover.append(node)
    main.append(cover)

    body = soup.new_tag("section", attrs={"class": "body"})
    for node in body_nodes:
        body.append(node)
    main.append(body)

    for heading in main.find_all("h2"):
        heading["class"] = ["chapter-title"]
    for table in main.find_all("table"):
        wrapper = soup.new_tag("div", attrs={"class": "table-wrap"})
        table.wrap(wrapper)
    for image in main.find_all("img"):
        image["class"] = ["hero-image"]

    css = r"""
@font-face { font-family: 'Noto Sans SC'; src: url('file:///home/joii/.local/share/fonts/NotoSansSC_400Regular.ttf') format('truetype'); font-weight: 400; }
@font-face { font-family: 'Noto Sans SC'; src: url('file:///home/joii/.local/share/fonts/NotoSansSC_500Medium.ttf') format('truetype'); font-weight: 500; }
@font-face { font-family: 'Noto Sans SC'; src: url('file:///home/joii/.local/share/fonts/NotoSansSC_700Bold.ttf') format('truetype'); font-weight: 700; }
@page { size: A4; margin: 18mm 17mm 18mm; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; color: #16384d; font-family: 'Noto Sans SC', sans-serif; font-size: 10.5pt; line-height: 1.75; }
body { background: #eef3f6; }
.document { width: 210mm; margin: 0 auto; background: #fff; }
.cover { min-height: 261mm; padding: 28mm 18mm 20mm; color: #fff; background: linear-gradient(145deg, #071a2f 0%, #0b3048 58%, #155671 100%); position: relative; overflow: hidden; break-after: page; }
.cover::before { content: ''; position: absolute; width: 145mm; height: 145mm; border-radius: 50%; right: -65mm; top: -55mm; background: rgba(11,187,214,.14); }
.cover::after { content: ''; position: absolute; width: 125mm; height: 125mm; border-radius: 50%; left: -60mm; bottom: -55mm; background: rgba(71,118,230,.13); }
.brand-mark { width: 24mm; height: 12mm; border-radius: 6mm; display: flex; align-items: center; justify-content: center; background: linear-gradient(90deg,#0bbbd6,#4776e6); font-weight: 700; letter-spacing: 1px; margin-bottom: 42mm; position: relative; z-index: 1; }
.cover h1 { font-size: 32pt; line-height: 1.22; margin: 0 0 8mm; letter-spacing: 1px; position: relative; z-index: 1; }
.cover blockquote { margin: 0 0 20mm; padding: 0 0 0 5mm; border-left: 1.6mm solid #2fd0e6; color: #ccebf3; font-size: 15pt; position: relative; z-index: 1; }
.cover p { max-width: 155mm; color: #d8e9ef; font-size: 11pt; position: relative; z-index: 1; }
.cover p:nth-of-type(1) { margin-top: 0; color: #9dd7e5; }
.body { padding: 0; }
.body > h2 { break-before: page; }
.body > h2:first-child { break-before: auto; }
h2 { color: #092b43; font-size: 22pt; line-height: 1.3; margin: 0 0 9mm; padding-top: 5mm; border-bottom: 1.2mm solid #d9eef3; padding-bottom: 4mm; }
h3 { display: block; width: 100%; clear: both; color: #0c6078; font-size: 14.5pt; line-height: 1.4; margin: 8mm 0 3mm; break-after: auto; page-break-after: auto; }
p { display: block; clear: both; margin: 0 0 4mm; text-align: justify; }
ul, ol { margin: 2mm 0 5mm 7mm; padding-left: 5mm; }
li { margin: 1.5mm 0; }
strong { color: #092b43; font-weight: 700; }
code { font-family: 'Noto Sans SC', sans-serif; color: #0b718c; background: #edf7fa; padding: .2mm 1mm; border-radius: 1mm; }
.hero-image { display: block; width: 100%; height: auto; margin: 4mm auto 6mm; border-radius: 3mm; break-inside: avoid; }
.table-wrap { width: 100%; margin: 4mm 0 6mm; break-inside: avoid; }
table { width: 100%; border-collapse: collapse; font-size: 9pt; line-height: 1.55; }
thead { display: table-header-group; }
tr { break-inside: avoid; }
th { color: #fff; background: #0b526d; font-weight: 600; text-align: left; padding: 2.5mm 2.6mm; border: .2mm solid #d2e4ea; }
td { vertical-align: top; padding: 2.4mm 2.6mm; border: .2mm solid #d2e4ea; }
tbody tr:nth-child(even) td { background: #f4f9fb; }
blockquote { margin: 4mm 0 6mm; padding: 4mm 5mm; background: #edf8fb; border-left: 1.4mm solid #0bbbd6; color: #315a6e; }
hr { border: 0; border-top: .4mm solid #dbe8ed; margin: 8mm 0; }
@media screen {
  .document { box-shadow: 0 0 30px rgba(9,43,67,.12); }
  .body { padding: 18mm 17mm; }
  .body > h2 { break-before: auto; margin-top: 12mm; }
}
@media print {
  body { background: #fff; }
  .document { width: auto; }
  .body { padding: 0; }
}
"""

    page = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>CPQ 系统白皮书</title>
  <style>{css}</style>
</head>
<body>{main}</body>
</html>
"""
    HTML_OUT.write_text(page, encoding="utf-8")


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, node) -> None:
    if isinstance(node, NavigableString):
        if str(node):
            set_run_font(paragraph.add_run(str(node)))
        return
    if not isinstance(node, Tag):
        return
    if node.name == "br":
        paragraph.add_run().add_break()
        return
    if node.name in {"strong", "b"}:
        run = paragraph.add_run(node.get_text())
        set_run_font(run, bold=True, color=NAVY)
        return
    if node.name in {"em", "i"}:
        run = paragraph.add_run(node.get_text())
        set_run_font(run)
        run.italic = True
        return
    if node.name == "code":
        run = paragraph.add_run(node.get_text())
        set_run_font(run, color="0B718C")
        return
    for child in node.children:
        add_inline(paragraph, child)


def add_rich_paragraph(document, tag: Tag, style=None, before=0, after=6):
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.45
    for child in tag.children:
        add_inline(paragraph, child)
    return paragraph


def add_table(document, tag: Tag) -> None:
    rows = tag.find_all("tr")
    if not rows:
        return
    col_count = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
    table = document.add_table(rows=len(rows), cols=col_count)
    table.autofit = True
    table.style = "Table Grid"
    for row_index, row_tag in enumerate(rows):
        cells = row_tag.find_all(["th", "td"], recursive=False)
        for col_index, cell_tag in enumerate(cells):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "0B526D")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F1F7F9")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.2
            for child in cell_tag.children:
                add_inline(paragraph, child)
            for run in paragraph.runs:
                set_run_font(run, size=8.5, bold=(row_index == 0), color=("FFFFFF" if row_index == 0 else None))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(document, tag: Tag) -> None:
    src = (ROOT / tag.get("src", "")).resolve()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run()
    run.add_picture(str(src), width=Inches(6.3))


def configure_doc_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("16384D")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.45

    for style_name, size, color in (("Title", 30, NAVY), ("Heading 1", 21, NAVY), ("Heading 2", 14, "0C6078")):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)

    if "Quote Box" not in styles:
        quote = styles.add_style("Quote Box", WD_STYLE_TYPE.PARAGRAPH)
    else:
        quote = styles["Quote Box"]
    quote.font.name = "Microsoft YaHei"
    quote._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    quote.font.size = Pt(11)
    quote.font.color.rgb = RGBColor.from_string("315A6E")
    quote.paragraph_format.left_indent = Cm(0.7)
    quote.paragraph_format.space_before = Pt(8)
    quote.paragraph_format.space_after = Pt(10)


def build_docx(source_soup: BeautifulSoup) -> None:
    document = Document()
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.different_first_page_header_footer = True
    configure_doc_styles(document)

    # Cover page
    mark = document.add_paragraph()
    mark.paragraph_format.space_before = Pt(16)
    run = mark.add_run("CPQ")
    set_run_font(run, size=16, bold=True, color=CYAN)
    mark.paragraph_format.space_after = Pt(110)

    first_hr = source_soup.find("hr")
    cover_nodes = []
    for node in source_soup.contents:
        if node is first_hr:
            break
        if isinstance(node, Tag):
            cover_nodes.append(node)

    for tag in cover_nodes:
        if tag.name == "h1":
            p = document.add_paragraph(style="Title")
            p.paragraph_format.space_after = Pt(16)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(tag.get_text())
            set_run_font(run, size=30, bold=True, color=NAVY)
        elif tag.name == "blockquote":
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(28)
            run = p.add_run(tag.get_text(" ", strip=True))
            set_run_font(run, size=15, color="0C748E")
        elif tag.name == "p":
            p = add_rich_paragraph(document, tag, after=9)
            for run in p.runs:
                set_run_font(run, size=10.5, color=GRAY)

    document.add_page_break()

    # Body
    body_started = False
    for tag in source_soup.contents:
        if tag is first_hr:
            body_started = True
            continue
        if not body_started or not isinstance(tag, Tag):
            continue
        if tag.name == "h2":
            if len(document.paragraphs) > 1 and tag.get_text(strip=True) != "目录":
                document.add_page_break()
            p = document.add_paragraph(style="Heading 1")
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(14)
            run = p.add_run(tag.get_text())
            set_run_font(run, size=21, bold=True, color=NAVY)
        elif tag.name == "h3":
            p = document.add_paragraph(style="Heading 2")
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(tag.get_text())
            set_run_font(run, size=14, bold=True, color="0C6078")
        elif tag.name == "p":
            image = tag.find("img", recursive=False)
            if image:
                add_image(document, image)
            else:
                add_rich_paragraph(document, tag)
        elif tag.name == "blockquote":
            add_rich_paragraph(document, tag, style="Quote Box", before=6, after=10)
        elif tag.name in {"ul", "ol"}:
            style = "List Bullet" if tag.name == "ul" else "List Number"
            for li in tag.find_all("li", recursive=False):
                add_rich_paragraph(document, li, style=style, after=3)
        elif tag.name == "table":
            add_table(document, tag)
        elif tag.name == "hr":
            document.add_paragraph()

    # Header and footer after the cover. Word uses one section, so keep them subtle.
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("CPQ 系统白皮书")
    set_run_font(header_run, size=8.5, color="7D96A5")
    add_page_number(section.footer.paragraphs[0])
    for run in section.footer.paragraphs[0].runs:
        set_run_font(run, size=8.5, color="7D96A5")

    core = document.core_properties
    core.title = "CPQ 系统白皮书"
    core.subject = "面向制造业的配置化报价、成本核算与决策协同平台"
    core.author = "CPQ 项目组"
    core.keywords = "CPQ, 制造业, 报价, 核价, 定价, 白皮书"
    document.save(DOCX_OUT)


def main() -> None:
    soup = render_markdown()
    build_html(soup)
    build_docx(soup)
    print(HTML_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    main()
